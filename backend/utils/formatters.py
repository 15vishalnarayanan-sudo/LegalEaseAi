from __future__ import annotations

from io import BytesIO
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt

from fpdf import FPDF

from backend.utils.text_utils import sanitize_text


# ============================================================
# COMMON HELPERS
# ============================================================

def _parse_lines(text: str) -> list[str]:
    cleaned = sanitize_text(text)

    return [line.strip() for line in cleaned.splitlines()]


def _clean_term(text: str) -> str:
    return re.sub(r"^\s*\d+[.)]\s*", "", text).strip()


def _is_numbered_section(line: str) -> bool:
    return bool(re.match(r"^\d+[.)]\s+", line))


def _is_heading(line: str) -> bool:
    return bool(line) and line.isupper() and len(line) <= 100


def _is_terms_heading(line: str) -> bool:
    upper = line.upper()

    return (
        "AGREED TERMS" in upper
        or "TERMS & CONDITIONS" in upper
    )


def _extract_terms(lines: list[str]) -> list[str]:
    """
    Extract terms from the section after
    AGREED TERMS or TERMS & CONDITIONS.
    """

    terms: list[str] = []
    inside_terms = False

    for line in lines:
        if not line:
            continue

        if _is_terms_heading(line):
            inside_terms = True
            continue

        if not inside_terms:
            continue

        # Stop only at known section headings.
        upper = line.upper()

        if (
            upper.endswith("GENERAL PROVISIONS")
            or upper.endswith("SIGNATURES")
            or upper.endswith("GENERAL")
            or upper.endswith("DECLARATIONS")
        ):
            break

        term = _clean_term(line)

        if term:
            terms.append(term)

    return terms


# ============================================================
# TXT
# ============================================================

def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


# ============================================================
# DOCX FONT
# ============================================================

def _set_docx_font(
    run,
    size: int = 11,
    bold: bool = False,
) -> None:

    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold


# ============================================================
# DOCX HEADING
# ============================================================

def _add_docx_heading(
    doc: Document,
    text: str,
    size: int = 11,
) -> None:

    paragraph = doc.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(9)
    paragraph.paragraph_format.space_after = Pt(4)

    run = paragraph.add_run(text)

    _set_docx_font(
        run,
        size=size,
        bold=True,
    )


# ============================================================
# DOCX TERMS TABLE
# ============================================================

def _add_terms_table(
    doc: Document,
    terms: list[str],
) -> None:

    if not terms:
        return

    table = doc.add_table(
        rows=1,
        cols=2,
    )

    table.style = "Table Grid"

    # Header
    header = table.rows[0].cells

    header[0].text = "S. No."
    header[1].text = "Terms & Conditions"

    for cell in header:

        cell.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

        for paragraph in cell.paragraphs:

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            for run in paragraph.runs:

                _set_docx_font(
                    run,
                    size=10,
                    bold=True,
                )

    # Rows
    for index, term in enumerate(
        terms,
        start=1,
    ):

        cells = table.add_row().cells

        cells[0].text = str(index)
        cells[1].text = term

        cells[0].vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

        cells[1].vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

        for paragraph in cells[0].paragraphs:

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            for run in paragraph.runs:

                _set_docx_font(
                    run,
                    size=10,
                )

        for paragraph in cells[1].paragraphs:

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
            )

            for run in paragraph.runs:

                _set_docx_font(
                    run,
                    size=10,
                )


# ============================================================
# DOCX FORMATTER
# ============================================================

def format_docx(
    text: str,
    doc_type: str,
    logo_bytes: bytes | None = None,
) -> bytes:

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    lines = _parse_lines(text)

    # --------------------------------------------------------
    # Logo
    # --------------------------------------------------------

    if logo_bytes:

        try:

            paragraph = doc.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run()

            run.add_picture(
                BytesIO(logo_bytes),
                width=Inches(1.25),
            )

        except Exception:
            pass

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    title = doc.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title.paragraph_format.space_after = Pt(10)

    title_run = title.add_run(
        sanitize_text(doc_type).upper()
    )

    _set_docx_font(
        title_run,
        size=16,
        bold=True,
    )

    # --------------------------------------------------------
    # Terms
    # --------------------------------------------------------

    terms = _extract_terms(lines)

    inside_terms = False

    # --------------------------------------------------------
    # Body
    # --------------------------------------------------------

    for line in lines:

        if not line:
            continue

        upper = line.upper()

        # ----------------------------------------------------
        # Terms heading
        # ----------------------------------------------------

        if _is_terms_heading(line):

            inside_terms = True

            _add_docx_heading(
                doc,
                line,
                size=11,
            )

            _add_terms_table(
                doc,
                terms,
            )

            continue

        # ----------------------------------------------------
        # Terms body
        # ----------------------------------------------------

        if inside_terms:

            if (
                upper.endswith("GENERAL PROVISIONS")
                or upper.endswith("SIGNATURES")
                or upper.endswith("GENERAL")
                or upper.endswith("DECLARATIONS")
            ):

                inside_terms = False

            else:

                continue

        # ----------------------------------------------------
        # Numbered heading
        # ----------------------------------------------------

        if _is_numbered_section(line):

            _add_docx_heading(
                doc,
                line,
                size=11,
            )

            continue

        # ----------------------------------------------------
        # All caps heading
        # ----------------------------------------------------

        if _is_heading(line):

            _add_docx_heading(
                doc,
                line,
                size=12,
            )

            continue

        # ----------------------------------------------------
        # Normal paragraph
        # ----------------------------------------------------

        paragraph = doc.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.15

        run = paragraph.add_run(line)

        _set_docx_font(
            run,
            size=11,
        )

    # --------------------------------------------------------
    # Footer
    # --------------------------------------------------------

    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        "LegalEase - AI-generated draft. "
        "Review with a qualified legal professional."
    )

    _set_docx_font(
        footer_run,
        size=8,
    )

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    output = BytesIO()

    doc.save(output)

    return output.getvalue()


# ============================================================
# PDF CLASS
# ============================================================

class LegalEasePDF(FPDF):

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            size=8,
        )

        self.cell(
            0,
            8,
            "LegalEase - AI-generated draft. Review before use.",
            align="C",
        )


# ============================================================
# PDF HELPERS
# ============================================================

def _safe_pdf_line(line: str) -> str:

    if not line:
        return ""

    line = line.replace(
        "\t",
        " ",
    )

    replacements = {
        "\u2013": "-",
        "\u2014": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2026": "...",
        "\u00a0": " ",
    }

    for old, new in replacements.items():

        line = line.replace(
            old,
            new,
        )

    line = re.sub(
        r"_{5,}",
        "____",
        line,
    )

    line = re.sub(
        r"\s+",
        " ",
        line,
    )

    return line.strip()


def _split_long_tokens(
    line: str,
    max_token_length: int = 35,
) -> str:

    tokens = line.split(" ")

    result: list[str] = []

    for token in tokens:

        if len(token) <= max_token_length:

            result.append(token)

        else:

            for index in range(
                0,
                len(token),
                max_token_length,
            ):

                result.append(
                    token[
                        index:index + max_token_length
                    ]
                )

    return " ".join(result)


def _write_pdf_line(
    pdf: FPDF,
    line: str,
    height: float = 6,
) -> None:

    if not line:

        pdf.ln(3)

        return

    line = _safe_pdf_line(line)

    if not line:

        pdf.ln(3)

        return

    line = _split_long_tokens(line)

    max_chars = 65

    chunks: list[str] = []

    while len(line) > max_chars:

        break_position = line.rfind(
            " ",
            0,
            max_chars,
        )

        if break_position <= 0:

            break_position = max_chars

        chunk = line[:break_position].strip()

        if chunk:
            chunks.append(chunk)

        line = line[break_position:].strip()

    if line:
        chunks.append(line)

    for chunk in chunks:

        pdf.cell(
            0,
            height,
            chunk,
            new_x="LMARGIN",
            new_y="NEXT",
        )


# ============================================================
# PDF FORMATTER
# ============================================================

def format_pdf(
    text: str,
    doc_type: str,
    logo_bytes: bytes | None = None,
) -> bytes:

    pdf = LegalEasePDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=18,
    )

    pdf.set_margins(
        left=15,
        top=15,
        right=15,
    )

    pdf.add_page()

    lines = _parse_lines(text)

    # --------------------------------------------------------
    # Logo
    # --------------------------------------------------------

    if logo_bytes:

        try:

            logo_stream = BytesIO(logo_bytes)

            pdf.image(
                logo_stream,
                x=(pdf.w - 30) / 2,
                y=10,
                w=30,
            )

            pdf.ln(28)

        except Exception:
            pass

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    pdf.set_font(
        "Helvetica",
        "B",
        16,
    )

    _write_pdf_line(
        pdf,
        sanitize_text(doc_type).upper(),
        10,
    )

    pdf.ln(4)

    # --------------------------------------------------------
    # Body
    # --------------------------------------------------------

    inside_terms = False

    for line in lines:

        if not line:

            pdf.ln(3)

            continue

        safe_line = _safe_pdf_line(line)

        if not safe_line:
            continue

        upper = safe_line.upper()

        # ----------------------------------------------------
        # Terms heading
        # ----------------------------------------------------

        if _is_terms_heading(safe_line):

            inside_terms = True

            pdf.set_font(
                "Helvetica",
                "B",
                11,
            )

            _write_pdf_line(
                pdf,
                safe_line,
                7,
            )

            continue

        # ----------------------------------------------------
        # Terms body
        # ----------------------------------------------------

        if inside_terms:

            if (
                upper.endswith("GENERAL PROVISIONS")
                or upper.endswith("SIGNATURES")
                or upper.endswith("GENERAL")
                or upper.endswith("DECLARATIONS")
            ):

                inside_terms = False

            else:

                term = _clean_term(safe_line)

                if term:

                    pdf.set_font(
                        "Helvetica",
                        size=10,
                    )

                    _write_pdf_line(
                        pdf,
                        f"- {term}",
                        6,
                    )

                continue

        # ----------------------------------------------------
        # All caps heading
        # ----------------------------------------------------

        if _is_heading(safe_line):

            pdf.set_font(
                "Helvetica",
                "B",
                11,
            )

            _write_pdf_line(
                pdf,
                safe_line,
                7,
            )

            continue

        # ----------------------------------------------------
        # Numbered heading
        # ----------------------------------------------------

        if _is_numbered_section(safe_line):

            pdf.set_font(
                "Helvetica",
                "B",
                10,
            )

            _write_pdf_line(
                pdf,
                safe_line,
                6,
            )

            continue

        # ----------------------------------------------------
        # Bullet
        # ----------------------------------------------------

        if safe_line.startswith("- "):

            pdf.set_font(
                "Helvetica",
                size=10,
            )

            _write_pdf_line(
                pdf,
                safe_line,
                6,
            )

            continue

        # ----------------------------------------------------
        # Normal text
        # ----------------------------------------------------

        pdf.set_font(
            "Helvetica",
            size=10,
        )

        _write_pdf_line(
            pdf,
            safe_line,
            6,
        )

    return bytes(pdf.output())


# ============================================================
# UNIVERSAL EXPORT
# ============================================================

def format_document(
    text: str,
    doc_type: str,
    output_format: str,
    logo_bytes: bytes | None = None,
) -> tuple[bytes, str, str]:

    output_format = output_format.lower().strip()

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    if output_format == "txt":

        return (
            format_txt(text),
            "text/plain; charset=utf-8",
            "legalease_document.txt",
        )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    if output_format == "docx":

        return (
            format_docx(
                text,
                doc_type,
                logo_bytes,
            ),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "legalease_document.docx",
        )

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if output_format == "pdf":

        return (
            format_pdf(
                text,
                doc_type,
                logo_bytes,
            ),
            "application/pdf",
            "legalease_document.pdf",
        )

    raise ValueError(
        "Unsupported export format."
    )